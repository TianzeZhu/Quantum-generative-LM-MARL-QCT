"""Generate the Word report from the scoped delivery data and figures."""
import argparse,json,sys
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
# For rebuilding, install python-docx. No model runtime is required.
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def main():
    out=ROOT/"03_documents"
    m=json.loads((out/"tables/experiment_summary.json").read_text(encoding="utf-8"))
    recon=json.loads((ROOT/"02_data/pretraining/reconstruction_metadata.json").read_text(encoding="utf-8"))
    d=Document()
    sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(2.4);sec.left_margin=sec.right_margin=Cm(2.5)
    sec.header_distance=sec.footer_distance=Cm(1.2)
    for name in ("Normal","Body Text"):
        st=d.styles[name];st.font.name="Times New Roman";st.font.size=Pt(12)
        st.element.rPr.rFonts.set(qn("w:eastAsia"),"宋体")
        st.paragraph_format.line_spacing=1.5;st.paragraph_format.space_after=Pt(5)
        st.paragraph_format.first_line_indent=Cm(.74)
    for name,size in (("Title",22),("Heading 1",15),("Heading 2",12)):
        st=d.styles[name];st.font.name="Times New Roman";st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.element.rPr.rFonts.set(qn("w:eastAsia"),"黑体")
        st.paragraph_format.first_line_indent=Cm(0);st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(8)
    d.styles["Caption"].font.size=Pt(10.5)
    d.styles["Caption"].element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"),"宋体")
    d.styles["Caption"].paragraph_format.first_line_indent=Cm(0)
    d.styles["Caption"].paragraph_format.line_spacing=1.1
    header=sec.header.paragraphs[0];header.text="串行量子序列模型的括号补全实验";header.alignment=WD_ALIGN_PARAGRAPH.CENTER
    header.paragraph_format.first_line_indent=Cm(0);header.runs[0].font.size=Pt(9)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.paragraph_format.first_line_indent=Cm(0)
    fld=OxmlElement("w:fldSimple");fld.set(qn("w:instr"),"PAGE");footer._p.append(fld)
    def para(text):
        p=d.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        return p
    def heading(text,level=1):d.add_heading(text,level=level)
    def eq(text,number):
        p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
        r=p.add_run(text+"    （"+str(number)+"）");r.font.name="Cambria Math";r.font.size=Pt(11)
    def table(caption,headers,rows,widths=None):
        p=d.add_paragraph(caption,style="Caption");p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        t=d.add_table(rows=1,cols=len(headers));t.autofit=False
        if widths:
            for c,w in zip(t.columns,widths):c.width=Cm(w)
        for i,s in enumerate(headers):t.rows[0].cells[i].text=str(s)
        for row in rows:
            cells=t.add_row().cells
            for c,s in zip(cells,row):c.text=str(s)
        for ri,row in enumerate(t.rows):
            pr=row._tr.get_or_add_trPr();n=OxmlElement("w:cantSplit");pr.append(n)
            if ri==0:
                repeat=OxmlElement("w:tblHeader");pr.append(repeat)
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next=ri<len(t.rows)-1
                    p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                    for rr in p.runs:rr.font.size=Pt(10);rr.bold=ri==0
                borders=OxmlElement("w:tcBorders")
                for edge in ("top","bottom"):
                    e=OxmlElement("w:"+edge)
                    e.set(qn("w:val"),"single" if ri==0 or (edge=="bottom" and ri==len(t.rows)-1) else "nil")
                    e.set(qn("w:sz"),"8");e.set(qn("w:color"),"000000");borders.append(e)
                cell._tc.get_or_add_tcPr().append(borders)
        d.add_paragraph().paragraph_format.space_after=Pt(0)
    def fig(name,caption,width=16):
        p=d.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(out/"figures"/name),width=Cm(width))
        p=d.add_paragraph(caption,style="Caption");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    page_count=0
    def page():
        nonlocal page_count
        page_count+=1
        if page_count<=2:d.add_page_break()
    p=d.add_paragraph();p.paragraph_format.space_after=Pt(65)
    p=d.add_paragraph("基于串行量子序列模型的\nDyck 括号补全实验",style="Title");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=d.add_paragraph("模型方法、两阶段训练与常规测试结果",style="Subtitle");p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
    for s in ("实验技术报告与可复现实验交付","实验配置：A1F1E3 · 单头 · 三量子比特 token","优化方案：0.003 → 0.001，总计 300 轮","实验标识：DYCK-A1F1E3-300","整理日期：2026 年 9 月 14 日"):
        p=d.add_paragraph(s);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
    page()
    heading("摘要")
    para("本报告研究串行量子序列到序列模型在短序列括号补全任务中的训练与测试表现。模型以三量子比特表示单个 token，结合可训练嵌入、旋转位置编码、基于 SwapTest 的成对注意力、经典概率聚合以及量子非线性与前馈变换，逐个生成目标 token。为降低模拟开销，内部工作辅助线以成功分支的 Kraus 映射处理，后续不再参与运算的寄存器通过偏迹约化；逻辑 pair 布局为 15 根量子线，解析模拟中最大的活动纯态为 7 根量子线。Softmax 与 Sigmoid 先在合成向量数据上独立进行保真度预训练，正式任务训练时冻结其参数。")
    para("实验采用四类括号构成的数据集，训练集、验证集与常规测试集分别包含 210、49 和 49 条样本。正式训练使用全序列交叉熵和 Adam 优化器，前 100 轮学习率为 0.003，后 200 轮为 0.001，保持参数及优化器状态连续，总计完成 63,000 次更新。按照验证 token 负对数似然选取的最佳模型位于第 298 轮。最终常规测试的 token 负对数似然为 0.1173，token 准确率为 96.69%，贪心生成整句准确率为 89.80%，即完整补全 44 条测试样本。实验结果说明该实现能够有效学习所设定短序列分布内的括号补全关系；结果属于经典计算机上的量子线路模拟，不构成量子硬件性能或量子计算优势的证据。")
    para("关键词：串行量子模型；序列到序列；括号补全；后选择；非线性预训练；两阶段优化")
    page()
    heading("1 引言")
    para("括号补全任务要求模型同时识别括号类型、配对关系以及未闭合结构的先后顺序，能够以较小的词表检验序列表示和条件生成过程。与一般文本生成不同，本任务具有确定的符号规则，正确输出可由后进先出栈直接计算，因此输入输出语义清晰，便于检查数据一致性以及模型预测是否正确。本报告以这一任务作为串行量子序列模型的功能性实验场景。")
    para("模型结构参考项目提供的 Quantum Generative Large Model Based on the Transformer Architecture 稿件[1]，并以本次归档源码[2]作为具体实现依据。本文关注成对注意力、概率级串行化、成功分支非线性变换以及端到端可微优化之间的组合关系。需要区分线路局部约化的等价性与整个序列处理方式的等价性：前者可在满足寄存器不再参与后续操作的条件下成立，后者涉及概率聚合及量子态重新制备，不能直接宣称与持续保留全序列相干信息的量子过程严格等价。")
    heading("2 数据集与任务定义")
    para("实验使用 DyckLanguage 合并数据中的训练、验证及常规测试划分。符号集合由圆括号、方括号、花括号和尖括号四种配对构成。每个输入是一段语法合法但尚未完全闭合的前缀，目标序列仅包含按后进先出顺序闭合剩余结构所需的右括号，不重复输入部分。例如，输入“< ( ) [”中圆括号已经完成配对，未闭合的部分依次为尖括号和方括号，因此正确输出为“] >”。")
    para("每个空白分隔的括号记作一个 token，不使用子词分词或额外的自然语言词表。输入长度与嵌套深度是不同概念：上述示例输入长度为 4，但同时未闭合括号的最大数量为 2，因此最大深度为 2。此次三个划分的输入长度均不超过 4，最大栈深度不超过 3；目标序列完整保留，不通过截断改变语义。数据由括号规则生成，记录的生成种子为 20260911，属于本项目的短序列任务配置。")
    table("表 1 数据集规模与序列长度（不含 BOS、EOS）",["划分","样本数","平均输入长度","平均输出长度","输入范围","输出范围"],[
        ["训练",210,"3.524","2.086","1–4","1–3"],["验证",49,"3.551","2.122","2–4","1–3"],["常规测试",49,"3.633","2.082","2–4","1–3"]],[2.3,1.6,3.1,3.1,2.3,2.3])
    para("训练集共含 740 个输入 token 和 438 个目标括号 token；加入每条样本的 EOS 后，每轮监督标签总数为 648。验证集与常规测试集分别包含 153 和 151 个含 EOS 的监督标签。三个划分之间的输入没有重复，验证与测试所需词汇均由训练词表覆盖。原始数据同时提供 JSONL 记录及逐行对齐的 .src/.tgt 文件，JSONL 保留输入、输出、序列长度、最大深度和生成来源等字段。")
    page()
    heading("3 串行量子模型")
    heading("3.1 总体计算结构",2)
    para("模型以目标位置为单位顺序计算。位置 t 的 query 在首个位置取 BOS，在后续位置取前一个目标 token；训练时该 token 来自真实目标序列，推理时来自已经生成的序列。每个位置的 context 包含全部源 token 与完整目标前缀。模型为每个 context token 分别构造成对量子运算，得到一个条件概率特征，再由宿主端聚合，执行一次 Softmax、第一前馈变换、Sigmoid 和第二前馈变换，最终输出当前目标位置的八维概率分布。不同位置共用同一组线路和嵌入参数，因此参数量不随序列长度增加。")
    fig("serial_architecture.png","图 1 串行量子序列模型的单目标位置计算流程")
    para("本实验采用单个模型 block，并非多层 block 堆叠。配置 A1F1E3 分别表示注意力 Q/K/V 中 SVD 酉部分的重复数为 1、前馈 GW1/GW2 中相应重复数为 1，以及 pair 混合纠缠模块的层数为 3。参数中的重复数描述模块内部的线路结构，不能等同于完整 Transformer 层数。JAX 实现通过扫描结构保持目标位置的串行依赖，同时允许相互独立的 context pair 在模拟器中向量化执行。")
    heading("3.2 嵌入与位置编码",2)
    para("源词表与目标词表彼此独立。源端包含 8 个括号符号和 4 个特殊 token，共 12 项；目标端包含 4 个右括号符号和 4 个特殊 token，共 8 项。数据文件采用的特殊 token 编号在加载时重新映射，模型宿主端固定为 PAD=0、BOS=1、EOS=2、UNK=3。宿主 token ID 用于查询嵌入表，不应与输出计算基编号混淆。")
    para("每个 token 的嵌入包含 7 个可训练实数，经有界非线性映射和补足归一化分量后构成八维实振幅态。设嵌入参数为 ℓ，前七个振幅取 aᵢ=(0.8/√7)tanh(ℓᵢ)，最后一个振幅取 √(1−Σᵢaᵢ²)，从而保证单位范数。源与目标嵌入合计包含 12×7+8×7=140 个可训练参数。随后在三根 token 量子线上分别施加 RX 旋转，角度为位置编号加一与共享位置参数的乘积。")
    eq("|eⱼ⟩ = [ ⊗ₖ RX((j+1)θₖ) ] |e(tokenⱼ)⟩",1)
    para("源位置从零开始编号，query 使用目标位置，作为 context 的目标前缀位置则以真实源长度为偏移。位置编码在 Q/K/V 变换前施加，三个位置信息参数跨所有 token 共享。训练中的 PAD 仅用于固定数组形状，通过 mask 排除其作为有效上下文或监督标签的贡献。")
    heading("3.3 成对注意力与后选择",2)
    para("对每个 context token，模型分别制备 query 对应的 Q 状态、经 V 变换的 query accumulator，以及 context 对应的 K 和 V 状态。Q/K 上执行完整 SwapTest 后，将 swap 辅助线、accumulator 和 context V 寄存器共同输入三层 StronglyEntanglingLayers。每层在七根线施加可训练单比特旋转，并按对应连接范围执行 CNOT 纠缠门。随后才取最终 swap=0 的联合概率行，而不是在纠缠层之前固定 swap 的测量结果。")
    eq("rⱼ(z)=P(Sfinal=0,A=z),   sⱼ=Σz rⱼ(z),   pⱼ(z)=rⱼ(z)/sⱼ",2)
    para("宿主端将成功分支概率转为注意力权重，并对条件概率特征加权求和。温度为一个共享可训练参数，其实现形式为 exp(clip(log τ,−5,5))；数值计算使用 ε=10⁻¹² 保护接近零的概率。温度调节的是不同 context pair 的聚合权重，不是输出词表的解码温度。")
    eq("αⱼ=softmaxⱼ[ln(sⱼ+ε)/τ],   pctx(z)=Σⱼ αⱼpⱼ(z)",3)
    para("聚合后以概率平方根重新制备 token 态，再进行非线性与前馈处理。该步骤仅保留概率特征，舍弃原状态的相对相位和跨 pair 相干关系，因此本文将其表述为概率特征级串行模型。它是当前实现的明确语义，而非将全部相干量子注意力不加区别地改写成经典求和。")
    heading("3.4 解析约化与线路宽度",2)
    para("固定逻辑 pair 布局包含四个 n 比特寄存器、一个 SwapTest 辅助比特及两根工作辅助线，宽度记为 4n+3。当 n=3 时，逻辑线宽为 15。SVD 内部的工作辅助线以测量零分支对应的 Kraus 算子直接作用于态向量，对角因子为 cos θ，并在保留成功分支后归一化。这种处理省去辅助线的显式枚举，但保留相应非酉条件映射。")
    eq("K₀=diag(cos θ₀,…,cos θ₇),   |ψ₀⟩=K₀|ψ⟩ / ‖K₀|ψ⟩‖₂",4)
    para("完整 SwapTest 后，若 K/Q 不再参与后续运算，可对这两个寄存器作无条件偏迹。swap 辅助线的两个分支权重为 (1+F)/2 和 (1−F)/2，其中 F=|⟨k|q⟩|²。解析模拟器分别推进这两个分支上的七线态，再混合其最终概率。两个分支都必须经过后续纠缠层，最终才选择 swap=0。K/Q 的偏迹不等于对 K/Q 作 postselect=0，也不应将最终成功概率与纠缠层之前的 SwapTest 零分支概率混为一谈。")
    para("在上述条件下，解析模拟的最大活动纯态宽度为 2n+1=7，对应 128 个复振幅。该数字描述单个局部状态，不代表 GPU 的总显存需求；向量化、自动微分、编译缓存和中间数组仍会占用额外资源。本实验在经典 GPU 上执行无有限 shots 抽样的概率计算，没有把后选择成功率折算为真实硬件重复执行次数。")
    heading("3.5 非线性、前馈与计算基读出",2)
    para("聚合态依次经过量子 Softmax、GW1、量子 Sigmoid 和 GW2。两个非线性模块使用相同的辅助线条件变换拓扑，但分别预训练独立参数。GW1/GW2 与 Q/K/V 均采用两个三比特酉块夹持可训练对角因子的结构，操作效果在成功分支归一化后输出。每个三比特酉块包含 63 个参数，单个 SVD 模块共有 2×63+8=134 个参数；这里的前馈处理保持八维 token 状态，不包含经典大词表输出头或额外的偏置向量。")
    para("最终 token 态按 Born 规则转为八个计算基的概率。有效右括号映射到基态 0–3，EOS 固定为基态 6，基态 4、5、7 视为 INVALID。不把无效输出概率剔除后重新归一化，因此流入无效基态的概率会降低正确标签的概率，并受到交叉熵目标的约束。")
    table("表 2 输出计算基与 token 的对应关系",["计算基编号","输出含义"],[[0,")"],[1,"]"],[2,"}"],[3,">"],[6,"EOS"],["4、5、7","INVALID"]],[5,11])
    page()
    heading("4 非线性函数预训练")
    para("非线性预训练使用在线合成的数值向量，而不是括号句子。对每个样本先在八维空间生成非负均匀随机向量，再进行 L2 归一化作为输入振幅。Softmax 目标为 softmax(βx) 经 L2 归一化得到的目标态，Sigmoid 目标为逐分量 sigmoid(x) 经 L2 归一化得到的目标态，β 固定为 1。函数作用对象是振幅向量；对应目标态的测量概率是目标振幅的模平方，不能直接等同于经典 Softmax 输出概率。")
    eq("|tsoft⟩=softmax(βx)/‖softmax(βx)‖₂,   |tsig⟩=sigmoid(x)/‖sigmoid(x)‖₂",5)
    para("预训练最小化输出态与目标态之间的一减平均保真度。对归一化纯态，保真度为内积模平方；该指标衡量量子态拟合程度，而不是括号预测准确率。Softmax 和 Sigmoid 各自独立执行 1600 次 Adam 更新，每批 16 个训练样本，并使用 256 个样本作最终验证。两组伪随机流程分别使用基准种子 2026 和偏移后的种子 2027，正式模型训练只读取最终参数并将其冻结。")
    eq("F=|⟨t|ψout⟩|²,   Lpre=1−(1/B)Σᵢ Fᵢ",6)
    table("表 3 非线性预训练配置与验证结果",["项目","Softmax","Sigmoid"],[
        ["数据维度","8","8"],["线路参数数","126","126"],["每批样本数","16","16"],["更新次数","1600","1600"],["训练向量总数","25,600","25,600"],["验证向量数","256","256"],["学习率","0.01","0.01"],["Adam β₁ / β₂","0.9 / 0.99","0.9 / 0.99"],["最终验证损失","0.015709","0.016914"],["最终验证保真度","0.984291","0.983086"]],[6,5,5])
    fig("nonlinear_pretraining.png","图 2 Softmax 与 Sigmoid 的预训练损失。纵轴为对数刻度，浅线保留原始更新值，实线为 31 步滑动平均，虚线表示最终验证损失。")
    para("归档的 paper.npz 与正式 checkpoint 记录的 SHA256 完全一致，确保交付的是本次实验实际加载的冻结参数，而非相同结构的其他权重。原始运行保存了参数、完整损失历史和日志，未逐批保存输入向量；交付时以相同 JAX 版本、随机种子和键分裂顺序重建训练及验证向量。重建文件明确标记为确定性重建数据，不标记为原始现场采样。重建验证样本计算出的两个损失与原始记录一致，重建过程、运行环境及文件摘要一并保存在 reconstruction_metadata.json。")
    page()
    heading("5 正式实验方案与配置")
    para("正式训练采用预先定义的两阶段学习率方案。第一阶段训练 100 轮，学习率设为 0.003；第二阶段在同一模型参数和 Adam 状态上训练至累计 300 轮，学习率降低为 0.001。阶段切换保留一阶动量、二阶动量和累计更新步数，仅改变学习率，其他超参数保持一致。两阶段分别执行 21,000 和 42,000 次更新，总计 63,000 次更新。")
    para("每条样本的监督标签为完整目标序列加 EOS。query 序列为 BOS 后接原目标序列，计算位置 t 时仅使用此前的目标 token 作为前缀。模型汇总该条样本全部有效标签位置的交叉熵，再做一次参数更新，不随机选择单个位置。batch_size 固定为 1，因而每条训练样本每轮对应一次更新。跨样本汇总时，应区分各序列等权的平均 CE 与按有效 token 数加权的 NLL。")
    eq("Lseq=−[1/(T+1)]Σₜ ln max{pθ(bₜ|x,y<t),ε}",7)
    para("优化器采用 Adam，β₁=0.9、β₂=0.99、εAdam=10⁻⁸，全局梯度范数上限为 5。每次更新后，嵌入参数和对数聚合温度裁剪到 [−5,5]，角参数投影到 [−π,π)。训练与验证使用完整数据，训练数据按由种子和轮次决定的顺序重新排列；每轮结束计算完整验证集指标，并按原始验证 token NLL 选择 best。latest 定期保存当前进度，best 与 latest 的用途分别是模型选择和恢复训练。")
    table("表 4 正式模型与优化配置",["配置项","取值"],[
        ["完整模型 block 数 / attention head 数","1 / 1"],["token 比特数 / 状态维度","3 / 8"],["A / F / E","1 / 1 / 3"],["逻辑 pair / 活动模拟线宽","15 / 7"],["可训练参数 / 冻结参数","877 / 252"],["学习率安排","第 1–100 轮：0.003；第 101–300 轮：0.001"],["batch_size / 验证 batch_size","1 / 1"],["随机种子","2026"],["梯度裁剪范数","5.0"],["每轮更新 / 总更新","210 / 63,000"],["每轮含 EOS 监督标签","648"],["训练数组 source / target_positions","4 / 4"],["周期 checkpoint","每 50 step；每轮末同时保存"],["最佳模型选择","验证 token NLL 最小"],["最大贪心解码长度","6 个位置"]],[6,10])
    table("表 5 可训练参数构成",["模块","参数数"],[
        ["源嵌入：12 × 7",84],["目标嵌入：8 × 7",56],["Q / K / V：3 × 134",402],["GW1 / GW2：2 × 134",268],["三层七线纠缠参数：3 × 7 × 3",63],["位置编码",3],["对数聚合温度",1],["合计",877]],[11,5])
    para("实际训练环境由日志记录为 Python 3.11.15、NumPy 2.2.6、JAX/JAXlib 0.7.1 和 PennyLane 0.44.1，使用 CUDA GPU 后端和双精度数值。训练快速路径是纯 JAX 解析模拟器，而不是运行在 PennyLane default.qubit 上的逐次 QNode 优化；PennyLane 参考线路用于结构表达和数值对照。部署脚本关闭显存预分配并要求 GPU 可用，无 GPU 时不静默回退。原始日志没有记录 GPU 型号和独立的逐样本预测文件，因此本文不据此补写这些信息。")
    page()
    heading("6 实验结果")
    heading("6.1 训练与验证过程",2)
    para("两阶段训练期间，训练与验证损失总体同步下降。图 3 展示各序列等权的交叉熵和按 token 加权的负对数似然，图 4 展示 token 与整句准确率。平滑曲线采用 11 轮居中滑动平均，两个学习率阶段分别计算，不跨阶段混合；原始曲线以浅线保留，同时另附未平滑图和逐轮 CSV。最佳 checkpoint 始终根据未平滑的验证 NLL 选取，图像平滑不会改变模型选择。")
    fig("training_loss_smoothed.png","图 3 两阶段训练与验证损失。虚线为学习率阶段边界；训练值为更新前在线统计，验证值为每轮固定 checkpoint 评估。")
    fig("training_accuracy_smoothed.png","图 4 两阶段训练准确率与验证准确率。token 指标包含 EOS，验证整句指标采用正确前缀。")
    para("验证 token NLL 最小值出现在第 298 轮、累计第 62,580 次更新，为 0.107167。该 checkpoint 的验证 token 准确率为 97.39%，验证整句准确率为 91.84%，即 49 条中有 45 条在正确前缀条件下完全正确。最终第 300 轮的验证 NLL 为 0.114503，略高于最佳值；因此正式测试选择第 298 轮 best，而不直接以最后一个 checkpoint 替代。")
    table("表 6 最佳与最终训练状态",["指标","最佳第 298 轮","最终第 300 轮"],[
        ["累计更新次数","62,580","63,000"],["训练在线 token NLL","0.094522","0.095582"],["训练在线 token 准确率","98.77%","98.46%"],["验证 token NLL","0.107167","0.114503"],["验证 token 准确率","97.39%","95.42%"],["验证正确前缀整句准确率","91.84%","85.71%"]],[7,4.5,4.5])
    para("训练指标是在一个 epoch 内、模型参数不断更新时逐样本记录并汇总的在线统计，不能将其表述为最佳 checkpoint 在完整训练集上的独立复测成绩。后期训练 token 准确率已接近 99%，说明模型能够较充分地拟合所设定任务。阶段后的总体改善同时包含更小学习率和更多训练步数的作用，仅凭本实验不能将二者贡献分别识别为因果效应。")
    heading("6.2 最终常规测试",2)
    para("最终评估使用 best checkpoint，对常规测试集全部 49 条样本计算正确前缀条件下的 NLL 和准确率，并进行使用模型自身输出作为后续输入的贪心解码。每条生成从 BOS 开始，选择最大概率基态作为当前输出，在生成 EOS 或 INVALID 时结束有效输出，最大有效生成长度设为 6。常规测试标签共包含 151 个含 EOS 的 token。")
    table("表 7 最终常规测试指标",["指标","结果"],[
        ["测试样本数","49"],["含 EOS 的目标 token 数","151"],["Token NLL","0.117266"],["序列等权平均 CE","0.130196"],["Token 准确率","96.69%（146/151）"],["正确前缀整句准确率","89.80%（44/49）"],["贪心生成整句准确率","89.80%（44/49）"],["EOS 终止率","100%"],["INVALID 率","0%"],["最大概率归一化误差","2.22 × 10⁻¹⁶"]],[8,8])
    para("贪心解码完整补全了 44 条测试样本，仍有 5 条未完全正确。正确前缀整句准确率与贪心整句准确率在本次结果中相同，但二者的定义不同，不能普遍互换。所有样本都正常生成 EOS，且没有 INVALID 输出，说明本次失败并非由无效基态或达到解码上限造成。由于原始评估未保存逐样本生成序列，不能进一步断言错误来自括号类型、输出顺序或提前结束中的哪一种情况。")
    heading("6.3 运行耗时与完整性",2)
    para("依据两次训练会话的开始和结束时间，第一阶段耗时约 3 小时 29 分 30 秒，第二阶段耗时约 5 小时 6 分 42 秒，累计训练含验证耗时约 8 小时 36 分 12 秒，统计时剔除了两阶段之间的停机间隔。最终常规测试内部计时为 200.31 秒，其中首次正确前缀评估约 71.86 秒，首次贪心解码约 125.80 秒。首次调用时间包含 JIT 编译与执行，不能据此当作稳定状态下单条样本的推理延迟。")
    para("归档检查确认训练日志包含连续的 63,000 条更新记录与 300 条 epoch 记录，汇总指标均为有限数值。best/latest 的文件摘要与各自元数据一致，预训练 artifact 的摘要也与训练 checkpoint 记录一致。训练末期每轮约 83 秒，其中验证首批次约 73 秒，提示验证调用仍存在较大的首次执行开销；当前代码每次验证新建 JIT 函数，这与重复编译的解释一致，但精确耗时归因仍需专门性能剖析。")
    page()
    heading("7 讨论")
    para("本实验在较小词表和确定性语法任务上验证了串行量子模型从数据编码、非线性参数预训练到完整序列优化及自回归解码的闭环。最终常规测试中较高的 token 与整句准确率表明，概率聚合后的固定维度 token 表示能够支持所设定短序列分布内的条件生成。与此同时，49 条测试样本的规模有限，单条样本对应约 2.04 个百分点的整句准确率变化，不能将单次结果视为对更广泛任务的充分统计证明。")
    para("对量子线路的解释需要保持边界清楚。SVD 工作辅助线的 Kraus 表达和后续闲置寄存器的偏迹，是在给定成功分支与后续操作范围内的局部解析约化；宿主端条件概率聚合与重新制备则引入了概率特征级处理。非线性模块的名称也不意味着其在全部输入态上严格等于经典函数，所报告的约 0.984 和 0.983 的验证保真度，仅对应既定合成输入分布。正式网络中位置旋转及酉变换可以产生复振幅，不能将该预训练验证结果不加限定地推广到所有中间态。")
    para("训练后期验证损失仍有小幅波动，选择最优 checkpoint 比机械采用最后一轮更符合既定评估流程。梯度裁剪在训练中较频繁地触发，说明优化仍受到条件归一化和参数化线路的数值特性影响，但有限梯度与持续改善的损失也表明训练并未因数值异常而中断。进一步工作应优先记录逐样本预测及置信度，并采用独立重复实验衡量结果稳定性；本报告不把可训练参数数量或单次准确率作为模型容量充分性的唯一依据。")
    heading("8 结论")
    para("本报告完成了 A1F1E3 单头串行量子序列模型在 Dyck 常规括号补全任务上的实验整理。模型以三量子比特 token 表示和 15 线逻辑 pair 结构实现逐 token 输出，利用 7 线解析活动态及可微宿主聚合降低模拟复杂度。Softmax 和 Sigmoid 经过独立保真度预训练后冻结，正式训练采用 0.003 与 0.001 的两阶段学习率及全序列交叉熵，共进行 300 轮优化。最终常规测试达到 96.69% 的 token 准确率与 89.80% 的贪心整句准确率，建立了可核验的数据、源码、参数与结果记录。")
    heading("参考文献与实验依据")
    for text in [
        "[1] Zhu T, Jiang X, Zhang L, Liao M. Quantum Generative Large Model Based on the Transformer Architecture. 项目提供的修订稿；不据此补写未核实的发表年份、期刊或 DOI。",
        "[2] 本项目串行模型源码快照：01_source/model/，含 serial_model_jax.py、statevector_jax_impl.py、nonlinear_pretrain_jax.py 及训练评估入口。文件摘要见交付清单。",
        "[3] 本次实验原始记录：02_data/pretraining/、02_data/training/ 与 02_data/evaluation/test_metrics.jsonl。指标以原始日志和 checkpoint 元数据为准。"]:
        p=para(text);p.paragraph_format.first_line_indent=Cm(0);p.runs[0].font.size=Pt(10.5)
    page()
    heading("附录 A 交付结构与复现说明")
    para("交付包由源代码、原始数据和实验文档三部分组成。01_source/model 保存交付时本地数值模型与训练实现快照，delivery_runner.py 提供相对路径启动入口，run_training.sh 和 run_evaluation.sh 分别执行两阶段训练及最终常规评估。数学核心代码未因交付而更改，相关源码的来源位置和摘要记录在 source_records.json；由于实验日志没有保存服务器 Git 提交号，源码归档不声称拥有未经记录的版本标识。")
    para("02_data/dataset/DyckLanguage 保存训练、验证和常规测试的原始 JSONL 与平行文本；原始样本与词表保持逐字节一致。交付 manifest 按本报告范围组织，delivery_index.json 保存包内数据摘要及原实验数据标识。可移植入口先校验包内文件，再使用原实验标识与 checkpoint 兼容；这一标识不等同于重新计算完整来源目录的摘要。该适配仅服务于数据归档与路径迁移，不改变模型计算或标签映射。")
    para("02_data/pretraining 保存实际冻结权重、配置、原始日志、损失历史及明确标识的重建输入输出样本；02_data/training 保存完整正式训练日志和 best/latest 两组原始参数及优化器状态；02_data/evaluation 仅保存最终常规测试记录。原始 checkpoint 中保留了历史服务器路径，交付入口显式使用包内预训练参数和数据路径，不要求重新创建原机器目录。")
    para("运行环境可依据 01_source/model/environment_h200.yml 创建，再通过 verify_delivery.py 核验归档。训练重现实验应使用新的输出目录，避免覆盖交付的原始日志与模型；运行 run_training.sh 可保持两阶段优化连续性，运行 run_evaluation.sh 默认读取包内 best 并评估常规测试。具体 shell 命令位于交付根目录 README.md。重新训练具有平台与依赖版本相关的数值差异，新的运行结果应单独保存，而不能替换为本报告的原始观测。")
    para("03_documents 包含本 Word 报告、PDF 阅读版、用于文档的 PNG/PDF 图像及逐轮统计表。图表可由 plot_delivery_figures.py 从归档记录重新生成，Word 文档由 build_word_report.py 生成，预训练数值样本可由 regenerate_pretraining.py 按记录配置重建。所有交付文件在最终清单中具有 SHA256 摘要，便于传输后检验一致性。")
    d.core_properties.title="基于串行量子序列模型的 Dyck 括号补全实验"
    d.core_properties.subject="A1F1E3两阶段训练与最终常规测试"
    d.core_properties.author="实验项目组"
    path=out/"串行量子模型_Dyck实验报告.docx"
    d.save(path)
    print(path)
if __name__=="__main__":main()
